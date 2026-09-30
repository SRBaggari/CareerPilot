"""Builders for real PDF/DOCX test documents (synthetic content only)."""

import io
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader, PdfWriter

SAMPLE_RESUME = (Path(__file__).parent / "fixtures" / "sample_resume.txt").read_text("utf-8")
SAMPLE_LINES = SAMPLE_RESUME.splitlines()


def _pdf_string(text: str) -> bytes:
    raw = text.encode("cp1252")
    return b"(" + raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)") + b")"


def make_pdf(lines: list[str], *, lines_per_page: int = 60, pages: int | None = None) -> bytes:
    """A text PDF (Helvetica, WinAnsi) with one text line per entry in ``lines``."""
    chunks = [lines[i : i + lines_per_page] for i in range(0, max(len(lines), 1), lines_per_page)]
    if pages is not None:
        chunks = (chunks * pages)[:pages]
    objects: list[bytes] = []
    page_ids = [3 + 2 * i for i in range(len(chunks))]
    font_id = 3 + 2 * len(chunks)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = b" ".join(b"%d 0 R" % pid for pid in page_ids)
    objects.append(b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(chunks)))
    for index, chunk in enumerate(chunks):
        body = (
            b"BT /F1 10 Tf 12 TL 50 760 Td "
            + b" ".join(_pdf_string(line) + b" Tj T*" for line in chunk)
            + b" ET"
        )
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>"
            % (font_id, page_ids[index] + 1)
        )
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(body), body))
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica "
                   b"/Encoding /WinAnsiEncoding >>")  # fmt: skip

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (number, obj))
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
              % (len(objects) + 1, xref))  # fmt: skip
    return out.getvalue()


def encrypt_pdf(pdf: bytes, password: str = "secret") -> bytes:  # noqa: S107 - test data
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(pdf)))
    writer.encrypt(password)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx(lines: list[str], *, table: list[list[str]] | None = None) -> bytes:
    """A DOCX where lines starting with a bullet become real Word list paragraphs."""
    document = Document()
    for line in lines:
        if line.startswith("\u2022 "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            document.add_paragraph(line)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_zip(files: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return out.getvalue()
